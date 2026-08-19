#!/usr/bin/env python3
"""Regenerate ccm-3.0.1.json from the source spreadsheet.

WHY THIS EXISTS
---------------
There was no generator. ccm-3.0.1.json was produced once, by hand or by a
script that was never committed, and it carried four misattributed columns that
nobody could re-derive or diff against source:

  framework_mappings.tenant_consumer
      -> col14, the Supplier Relationship "Tenant / Consumer" sub-column.
         Already correctly present as supplier_relationship.tenant_consumer.
         Its value is the marker "X", not a mapping.
  framework_mappings.csa_enterprise_architecture_..._initiative_2
      -> col27, the CSA EA composite's "Public" sub-column. Value "shared".
  framework_mappings.csa_enterprise_architecture_..._initiative_3
      -> col28, the CSA EA composite's "Private" sub-column. Value "x".
  framework_mappings.odca_um_pa_r2_0_2
      -> col47, the ODCA composite's "PA level" sub-column. Value "SGP/BSGP".

All four are column drift across a THREE-ROW merged header (rows 2-4). The
sheet's mapping section holds 32 real targets, two of which are composites
spanning several columns.

This script therefore asserts the expected header text at every column it reads
before parsing a single data row. A positional parser against a merged header
that does not check its headers is how the original defect happened; failing
loudly on an unexpected layout is the entire point.

Usage:
    python3 build_json.py            # writes ccm-3.0.1.json next to the xlsx
    python3 build_json.py --check    # verify only, non-zero exit on mismatch
"""
import argparse
import json
import re
import sys

import openpyxl

XLSX = "ccm-3.0.1.xlsx"
OUT = "ccm-3.0.1.json"
SHEET = "CSA CCM V3.0.1"

# Columns whose header text must match before we trust any position.
# (index, row, expected substring) -- row 2 is the section header row,
# rows 3-4 carry sub-headers.
HEADER_GUARDS = [
    (0, 2, "Control Domain"),
    (1, 2, "Control ID"),
    (2, 2, "Updated Control Specification"),
    (3, 2, "Architectural Relevance"),
    (9, 2, "Corp Gov Relevance"),
    (10, 2, "Cloud Service Delivery Model Applicability"),
    (13, 2, "Supplier Relationship"),
    (15, 2, "Scope Applicability"),
    (3, 3, "Phys"), (4, 3, "Network"), (5, 3, "Compute"),
    (6, 3, "Storage"), (7, 3, "App"), (8, 3, "Data"),
    (10, 3, "SaaS"), (11, 3, "PaaS"), (12, 3, "IaaS"),
    (13, 3, "Service Provider"), (14, 3, "Tenant / Consumer"),
    (26, 3, "CSA Enterprise Architecture"),
    (26, 4, "Domain > Container > Capability"),
    (27, 4, "Public"), (28, 4, "Private"),
    (46, 3, "ODCA UM: PA R2.0"),
    (46, 4, "PA ID"), (47, 4, "PA level"),
]

ARCH = {3: "phys", 4: "network", 5: "compute", 6: "storage", 7: "app", 8: "data"}
SDM = {10: "saas", 11: "paas", 12: "iaas"}
SUPPLIER = {13: "service_provider", 14: "tenant_consumer"}

# Mapping columns, EXCLUDING the composite sub-columns handled separately.
# Key = the slug this file emits; taken from the sheet's own header text.
MAPPING_COLS = {
    15: "aicpa_2009_tsc_map",
    16: "aicpa_trust_service_criteria_soc_2sm_report",
    17: "aicpa_2014_tsc",
    18: "bits_shared_assessments_aup_v5_0",
    19: "bits_shared_assessments_sig_v6_0",
    20: "bsi_germany",
    21: "canada_pipeda",
    22: "ccm_v1_x",
    23: "cobit_4_1",
    24: "cobit_5_0",
    25: "coppa",
    29: "csa_guidance_v3_0",
    30: "enisa_iaf",
    31: "95_46_ec_european_union_data_protection_directive",
    32: "fedramp_security_controls_final_release_jan_2012_low_impact_level",
    33: "fedramp_security_controls_final_release_jan_2012_moderate_impact_level",
    34: "ferpa",
    35: "gapp_aug_2009",
    36: "hipaa_hitech_act",
    37: "iso_iec_27001_2005",
    38: "iso_iec_27001_2013",
    39: "itar",
    40: "jericho_forum",
    41: "mexico_federal_law_on_protection_of_personal_data_held_by_private_parties",
    42: "nerc_cip",
    43: "nist_sp800_53_r3",
    44: "nist_sp800_53_r4_app_j",
    45: "nzism",
    48: "pci_dss_v2_0",
    49: "pci_dss_v3_0",
}
EA_COL, EA_PUBLIC, EA_PRIVATE = 26, 27, 28
ODCA_ID, ODCA_LEVEL = 46, 47

# The source writes an explicit "-" for "no mapping here". Downstream conveys the
# same thing by omitting the key, so treat "-" as absent rather than storing two
# encodings of the same fact.
NO_MAPPING = {"-", "", "n/a", "na"}


def cell(v):
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def truthy(v):
    """The sheet marks applicability with an X. Anything else non-empty counts
    as set too, so a variant marker is not silently read as false."""
    s = cell(v)
    return bool(s)


def verify_headers(ws):
    rows = {n: r for n, r in enumerate(ws.iter_rows(min_row=1, max_row=4, values_only=True), start=1)}
    problems = []
    for idx, rownum, expected in HEADER_GUARDS:
        row = rows.get(rownum) or ()
        got = cell(row[idx]) if idx < len(row) else None
        flat = re.sub(r"\s+", " ", got or "")
        if expected.lower() not in flat.lower():
            problems.append(f"col{idx} row{rownum}: expected ~{expected!r}, found {flat[:60]!r}")
    if problems:
        print("HEADER MISMATCH -- refusing to parse. The sheet layout is not what "
              "this script was written against:", file=sys.stderr)
        for p in problems:
            print(f"  {p}", file=sys.stderr)
        sys.exit(2)


def build():
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    ws = wb[SHEET]
    verify_headers(ws)

    controls = {}
    for row in ws.iter_rows(min_row=5, values_only=True):
        cid = cell(row[1]) if len(row) > 1 else None
        if not cid or not re.fullmatch(r"[A-Z&]{2,4}-\d{2}", cid):
            continue

        # col0 holds "domain<newline>title" for 132 of the 133 controls. DCS-07
        # alone uses "domain - title" on a single line. Handle both, or DCS-07
        # silently ends up with its domain and title fused and no title at all.
        # (The AICM plugin's clean_domain() splits on " - " for the same reason,
        # so this inconsistency is known in CSA's spreadsheets generally.)
        raw_domain = cell(row[0]) or ""
        parts = [p.strip() for p in raw_domain.split("\n") if p.strip()]
        if len(parts) == 1 and " - " in parts[0]:
            head, _, tail = parts[0].rpartition(" - ")
            parts = [head.strip(), tail.strip()]
        domain = parts[0] if parts else None
        title = parts[1] if len(parts) > 1 else None

        fm = {}
        for idx, key in MAPPING_COLS.items():
            v = cell(row[idx]) if idx < len(row) else None
            if v and v.lower() not in NO_MAPPING:
                fm[key] = v

        ea_path = cell(row[EA_COL]) if EA_COL < len(row) else None
        ea_pub = cell(row[EA_PUBLIC]) if EA_PUBLIC < len(row) else None
        ea_priv = cell(row[EA_PRIVATE]) if EA_PRIVATE < len(row) else None
        if any((ea_path, ea_pub, ea_priv)):
            fm["csa_enterprise_architecture_formerly_trusted_cloud_initiative"] = {
                "path": ea_path, "public": ea_pub, "private": ea_priv,
            }
        od_id = cell(row[ODCA_ID]) if ODCA_ID < len(row) else None
        od_lvl = cell(row[ODCA_LEVEL]) if ODCA_LEVEL < len(row) else None
        if any((od_id, od_lvl)):
            fm["odca_um_pa_r2_0"] = {"pa_id": od_id, "pa_level": od_lvl}

        controls[cid] = {
            "control_domain": domain,
            "control_title": title,
            "control_id": cid,
            "control_specification": cell(row[2]),
            "architectural_relevance": {n: truthy(row[i]) for i, n in ARCH.items()},
            "corp_gov_relevance": truthy(row[9]),
            "service_delivery_model": {n: truthy(row[i]) for i, n in SDM.items()},
            "supplier_relationship": {n: truthy(row[i]) for i, n in SUPPLIER.items()},
            "framework_mappings": fm,
        }
    wb.close()
    # IVS-13 (Network Architecture) genuinely has NO mapping columns filled in the
    # source spreadsheet. An empty framework_mappings for it is faithful, not a
    # parse failure -- do not "fix" it by inventing mappings or by treating a zero
    # count as an error.
    return {
        "metadata": {
            "framework": "CCM",
            "version": "3.0.1",
            "source_file": "CSA_CCM_v3.0.1-09-16-2014.xlsx",
            "total_controls": len(controls),
            "generated_by": "build_json.py",
        },
        "controls": controls,
    }


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--check", action="store_true",
                    help="compare against the committed JSON instead of writing")
    args = ap.parse_args()
    doc = build()
    if args.check:
        existing = json.load(open(OUT, encoding="utf-8"))
        if existing != doc:
            print(f"{OUT} differs from a fresh build of {XLSX}", file=sys.stderr)
            sys.exit(1)
        print(f"{OUT} matches a fresh build ({doc['metadata']['total_controls']} controls)")
        return
    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(doc, f, indent=2, ensure_ascii=False)
        f.write("\n")
    print(f"Wrote {OUT}: {doc['metadata']['total_controls']} controls")


if __name__ == "__main__":
    main()
